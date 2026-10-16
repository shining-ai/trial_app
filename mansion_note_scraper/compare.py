"""
全 run_dir の results.csv を統合し、同一物件の時系列推移を比較する Excel を生成する。
写真は各物件の最新 run_dir から取得して日付列の右側に埋め込む。

使い方:
    docker compose run --rm compare
    docker compose run --rm compare --last 5   # 直近5回分のみ
    docker compose run --rm compare --output /app/output/my_comparison.xlsx
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

OUTPUT_ROOT = Path("/app/output")

# 静的情報（run をまたいで基本的に変わらない列）
STATIC_FIELDS: list[tuple[str, str, int]] = [
    # (field_name, header_label, column_width)
    ("building_name", "物件名",   20),
    ("address",       "住所",     20),
    ("layout",        "間取り",   12),
    ("area",          "面積",      8),
    ("floor",         "所在階",   10),
    ("direction",     "向き",      8),
    ("built_year",    "築年月",   12),
    ("station",       "交通",     30),
    ("url",           "URL",      40),
]

# 時系列比較する動的フィールド（賃料＋管理費は合算して1列）
DYNAMIC_FIELDS: list[tuple[str, str]] = [
    ("total_rent", "合計賃料\n(賃料+管理費)"),
    ("shikikin",   "敷金"),
    ("reikin",     "礼金"),
]

# 写真セルの表示サイズ（to_excel.py と同じ設定）
PHOTO_WIDTH_PX  = 400
PHOTO_HEIGHT_PX = 450
ROW_HEIGHT      = 338   # points (450px ÷ 1.333)
PHOTO_COL_WIDTH = 57    # characters (400px ÷ 7)

# 特定カテゴリ列（to_excel.py と同じ）
SPECIAL_COLUMNS: list[tuple[str, list[str]]] = [
    ("間取り図", ["間取り"]),
    ("キッチン", ["キッチン", "台所", "コンロ"]),
    ("トイレ",   ["トイレ"]),
    ("浴室",     ["浴室", "バスルーム", "バス", "お風呂"]),
    ("洗面台",   ["洗面"]),
]

# カラーパレット（run ごとに交互）
_DATE_BG = ["1F4E79", "2A5585"]

ABSENT_FILL  = PatternFill("solid", fgColor="DDDDDD")   # 掲載なし（グレー）
UP_FILL      = PatternFill("solid", fgColor="FFB3B3")   # 値上がり（薄赤）
DOWN_FILL    = PatternFill("solid", fgColor="B3CFFF")   # 値下がり（薄青）
NEW_FILL     = PatternFill("solid", fgColor="B3FFD9")   # 新着（薄緑）
GONE_FILL    = PatternFill("solid", fgColor="FFDDDD")   # 消滅（薄ピンク）


# ---------------------------------------------------------------------------
# テキスト系ユーティリティ
# ---------------------------------------------------------------------------

def _parse_rent(s: str) -> int | None:
    m = re.search(r"([\d.]+)万", s)
    return int(float(m.group(1)) * 10000) if m else None


def _parse_kanrihi(s: str) -> int | None:
    cleaned = re.sub(r"[,円\s]", "", s)
    if not cleaned or cleaned in ("-", "なし", "無料"):
        return 0
    return int(cleaned) if re.fullmatch(r"\d+", cleaned) else None


def _fmt_yen(val: int | None, show_sign: bool = False) -> str:
    if val is None:
        return ""
    sign = "+" if show_sign and val > 0 else ""
    if abs(val) >= 10000:
        return f"{sign}{val / 10000:.1f}万円"
    return f"{sign}{val}円"


def _run_label(run_dir: Path) -> str:
    try:
        return datetime.strptime(run_dir.name[:15], "%Y%m%d_%H%M%S").strftime("%Y/%m/%d %H:%M")
    except ValueError:
        return run_dir.name


def _label_to_dt(label: str) -> datetime | None:
    """_run_label が生成したラベルを datetime に戻す（不正な形式は None）。"""
    try:
        return datetime.strptime(label, "%Y/%m/%d %H:%M")
    except ValueError:
        return None


def _find_all_run_dirs(output_root: Path) -> list[Path]:
    return sorted(
        d for d in output_root.iterdir()
        if d.is_dir() and (d / "results.csv").exists()
    )


def _load_run(run_dir: Path) -> dict[str, dict]:
    result: dict[str, dict] = {}
    with open(run_dir / "results.csv", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            hid = row.get("house_id", "")
            if hid:
                result[hid] = row
    return result


# ---------------------------------------------------------------------------
# 写真系ユーティリティ（to_excel.py と同じロジック）
# ---------------------------------------------------------------------------

def _find_photo_dir(photos_dir: Path, house_id: str) -> Path | None:
    if not photos_dir.exists():
        return None
    for d in photos_dir.iterdir():
        if d.is_dir() and d.name.endswith(f"_{house_id}"):
            return d
    return None


def _find_all_photos(photos_dir: Path, house_id: str) -> list[Path]:
    d = _find_photo_dir(photos_dir, house_id)
    if d is None:
        return []
    return sorted(
        [p for p in d.iterdir() if p.suffix.lower() in (".jpg", ".jpeg", ".png", ".webp")],
        key=lambda p: p.stem,
    )


def _load_labels(photos_dir: Path, house_id: str) -> dict[str, str]:
    d = _find_photo_dir(photos_dir, house_id)
    if d is None:
        return {}
    label_file = d / "labels.json"
    if not label_file.exists():
        return {}
    try:
        return json.loads(label_file.read_text(encoding="utf-8"))
    except Exception:
        return {}


def _match_special(label: str) -> str | None:
    for col_name, keywords in SPECIAL_COLUMNS:
        if any(kw in label for kw in keywords):
            return col_name
    return None


def _scale_image(img_path: Path, max_w: int, max_h: int) -> tuple[int, int]:
    with PILImage.open(img_path) as im:
        orig_w, orig_h = im.size
    ratio = min(max_w / orig_w, max_h / orig_h)
    return int(orig_w * ratio), int(orig_h * ratio)


def _embed_photo(ws, photo_path: Path, cell_addr: str) -> None:
    ws[cell_addr].alignment = Alignment(horizontal="center", vertical="center")
    try:
        w, h = _scale_image(photo_path, PHOTO_WIDTH_PX, PHOTO_HEIGHT_PX)
        with PILImage.open(photo_path) as pil_img:
            pil_img = pil_img.convert("RGB")
            buf = io.BytesIO()
            pil_img.save(buf, format="JPEG")
            buf.seek(0)
        img = XLImage(buf)
        img.width  = w
        img.height = h
        ws.add_image(img, cell_addr)
    except Exception as e:
        ws[cell_addr] = "読込失敗"
        print(f"  ⚠ 写真埋め込み失敗 ({photo_path.name}): {e}")


def _find_latest_run_photos(
    run_dirs: list[Path], hid: str
) -> tuple[list[Path], dict[str, str]]:
    """最新 run_dir から遡って写真を探し、最初に見つかった run の写真を返す。"""
    for rd in reversed(run_dirs):
        photos_dir = rd / "photos"
        photos = _find_all_photos(photos_dir, hid)
        if photos:
            return photos, _load_labels(photos_dir, hid)
    return [], {}


# ---------------------------------------------------------------------------
# Sheet 1: 時系列推移マトリクス
# ---------------------------------------------------------------------------

def _build_matrix_sheet(
    wb: Workbook,
    runs: list[tuple[str, dict[str, dict]]],
    run_dirs: list[Path],
) -> None:
    ws = wb.active
    ws.title = "時系列推移"

    n_static  = len(STATIC_FIELDS)
    n_dynamic = len(DYNAMIC_FIELDS)

    static_fill  = PatternFill("solid", fgColor="1F4E79")
    static_font  = Font(color="FFFFFF", bold=True)
    special_fill = PatternFill("solid", fgColor="2E75B6")
    special_font = Font(color="FFFFFF", bold=True)

    # ---- 全 house_id を収集（ヘッダー構築前に写真列数が必要なため先に計算） ----
    all_ids: set[str] = set()
    for _, data in runs:
        all_ids.update(data.keys())

    first_data: dict[str, dict] = {}
    for hid in all_ids:
        for _, data in runs:
            if hid in data:
                first_data[hid] = data[hid]
                break

    sorted_ids = sorted(
        all_ids,
        key=lambda hid: (first_data[hid].get("building_name", ""), hid),
    )

    # 写真を事前収集してその他写真の最大数を算出
    photo_cache: dict[str, tuple[list[Path], dict[str, str]]] = {}
    for hid in all_ids:
        photo_cache[hid] = _find_latest_run_photos(run_dirs, hid)

    max_other_photos = 0
    for hid in all_ids:
        photos, labels = photo_cache[hid]
        used: set[Path] = set()
        for _, keywords in SPECIAL_COLUMNS:
            matched = next(
                (p for p in photos if p not in used
                 and any(kw in labels.get(p.name, "") for kw in keywords)),
                None,
            )
            if matched:
                used.add(matched)
        max_other_photos = max(max_other_photos, len([p for p in photos if p not in used]))

    # 写真列の開始列
    photo_special_start = n_static + 1 + len(runs) * n_dynamic
    photo_other_start   = photo_special_start + len(SPECIAL_COLUMNS)

    # ---- ヘッダー行 1〜2 ----
    # 静的列: 行1〜2 をマージ
    for c, (_, label, width) in enumerate(STATIC_FIELDS, start=1):
        ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
        cell = ws.cell(1, c, label)
        cell.fill      = static_fill
        cell.font      = static_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width

    # 動的列: run ごとにグループヘッダー（行1）+ フィールド名（行2）
    for r, (label, _) in enumerate(runs):
        start_c = n_static + 1 + r * n_dynamic
        end_c   = start_c + n_dynamic - 1
        fill    = PatternFill("solid", fgColor=_DATE_BG[r % len(_DATE_BG)])
        font    = Font(color="FFFFFF", bold=True)

        ws.merge_cells(start_row=1, start_column=start_c, end_row=1, end_column=end_c)
        cell = ws.cell(1, start_c, label)
        cell.fill      = fill
        cell.font      = font
        cell.alignment = Alignment(horizontal="center", vertical="center")

        for d, (_, d_label) in enumerate(DYNAMIC_FIELDS):
            cell = ws.cell(2, start_c + d, d_label)
            cell.fill      = fill
            cell.font      = font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(start_c + d)].width = 10

    # 写真列ヘッダー: 特定カテゴリ（行1〜2 マージ）
    for i, (col_name, _) in enumerate(SPECIAL_COLUMNS):
        col = photo_special_start + i
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
        cell = ws.cell(1, col, col_name)
        cell.fill      = special_fill
        cell.font      = special_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = PHOTO_COL_WIDTH

    # 写真列ヘッダー: その他（行1〜2 マージ）
    for i in range(max_other_photos):
        col = photo_other_start + i
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
        cell = ws.cell(1, col, f"写真{i + 1}")
        cell.fill      = static_fill
        cell.font      = static_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = PHOTO_COL_WIDTH

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18

    # ---- データ行 ----
    for row_n, hid in enumerate(sorted_ids, start=3):
        ws.row_dimensions[row_n].height = ROW_HEIGHT
        base = first_data[hid]

        # 静的情報
        for c, (field, _, _) in enumerate(STATIC_FIELDS, start=1):
            cell = ws.cell(row_n, c, base.get(field, ""))
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=field in ("station", "building_name", "address"),
            )

        # 動的情報（前 run との比較で色付け）
        prev: dict[str, int | None] = {}
        for r, (_, data) in enumerate(runs):
            start_c = n_static + 1 + r * n_dynamic

            if hid not in data:
                for d in range(n_dynamic):
                    cell = ws.cell(row_n, start_c + d, "")
                    cell.fill      = ABSENT_FILL
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                continue

            row = data[hid]
            rent_val    = _parse_rent(row.get("rent", ""))
            kanrihi_val = _parse_kanrihi(row.get("kanrihi", ""))
            total_rent  = (rent_val + (kanrihi_val or 0)) if rent_val is not None else None
            parsed = {
                "total_rent": total_rent,
                "shikikin":   row.get("shikikin", ""),
                "reikin":     row.get("reikin", ""),
            }

            for d, (field, _) in enumerate(DYNAMIC_FIELDS):
                col = start_c + d
                val = parsed[field]
                if field == "total_rent":
                    display = _fmt_yen(val) if isinstance(val, int) else row.get("rent", "")
                else:
                    display = row.get(field, "")
                cell = ws.cell(row_n, col, display)
                cell.alignment = Alignment(
                    horizontal="right" if isinstance(val, int) else "left",
                    vertical="center",
                )
                if field == "total_rent" and isinstance(val, int) and isinstance(prev.get(field), int):
                    if val > prev[field]:       # type: ignore[operator]
                        cell.fill = UP_FILL
                    elif val < prev[field]:     # type: ignore[operator]
                        cell.fill = DOWN_FILL
                if field == "total_rent" and isinstance(val, int):
                    prev[field] = val

        # 写真の埋め込み
        photos, labels = photo_cache[hid]
        used_photos: set[Path] = set()

        for i, (_, keywords) in enumerate(SPECIAL_COLUMNS):
            col       = photo_special_start + i
            cell_addr = f"{get_column_letter(col)}{row_n}"
            matched   = None
            for kw in keywords:
                matched = next(
                    (p for p in photos if p not in used_photos and kw in labels.get(p.name, "")),
                    None,
                )
                if matched:
                    break
            if matched:
                used_photos.add(matched)
                _embed_photo(ws, matched, cell_addr)

        other_photos = [p for p in photos if p not in used_photos]
        for j, photo_path in enumerate(other_photos):
            col       = photo_other_start + j
            cell_addr = f"{get_column_letter(col)}{row_n}"
            _embed_photo(ws, photo_path, cell_addr)

        if row_n % 20 == 2:
            print(f"  {row_n - 2} 件完了...")

    ws.freeze_panes = "A3"
    total_photo_cols = len(SPECIAL_COLUMNS) + max_other_photos
    print(f"  物件数: {len(sorted_ids)} 件, 写真列: {total_photo_cols} 列")


# ---------------------------------------------------------------------------
# Sheet 2: 変動サマリー（新着・消滅・賃料変動）
# ---------------------------------------------------------------------------

def _build_diff_sheet(wb: Workbook, runs: list[tuple[str, dict[str, dict]]]) -> None:
    ws = wb.create_sheet("変動サマリー")

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    HEADERS = [
        ("変動種別", 12), ("日付",          14), ("物件名",        20),
        ("住所",     20), ("間取り",        12), ("面積",           8),
        ("交通",     30), ("変動前 合計賃料", 14), ("変動後 合計賃料", 14),
        ("差額",     10), ("URL",           40),
    ]
    for c, (label, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(1, c, label)
        cell.fill      = header_fill
        cell.font      = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width

    row_n = 2
    prev_data: dict[str, dict] | None = None

    for label, data in runs:
        if prev_data is None:
            prev_data = data
            continue

        prev_ids = set(prev_data.keys())
        curr_ids = set(data.keys())

        def _write_row(values: list, fill: PatternFill) -> None:
            nonlocal row_n
            WRAP_COLS = {3, 4, 7, 11}
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row_n, c, v)
                cell.fill      = fill
                cell.alignment = Alignment(vertical="center", wrap_text=(c in WRAP_COLS))
            row_n += 1

        def _total(row: dict) -> int | None:
            rent_val    = _parse_rent(row.get("rent", ""))
            kanrihi_val = _parse_kanrihi(row.get("kanrihi", ""))
            return (rent_val + (kanrihi_val or 0)) if rent_val is not None else None

        for hid in sorted(curr_ids - prev_ids, key=lambda h: data[h].get("building_name", "")):
            r = data[hid]
            _write_row([
                "新着", label,
                r.get("building_name", ""), r.get("address", ""),
                r.get("layout", ""),        r.get("area", ""),
                r.get("station", ""),
                "", _fmt_yen(_total(r)), "",
                r.get("url", ""),
            ], NEW_FILL)

        for hid in sorted(prev_ids - curr_ids, key=lambda h: prev_data[h].get("building_name", "")):  # type: ignore[index]
            r = prev_data[hid]  # type: ignore[index]
            _write_row([
                "消滅", label,
                r.get("building_name", ""), r.get("address", ""),
                r.get("layout", ""),        r.get("area", ""),
                r.get("station", ""),
                _fmt_yen(_total(r)), "", "",
                r.get("url", ""),
            ], GONE_FILL)

        for hid in sorted(curr_ids & prev_ids, key=lambda h: data[h].get("building_name", "")):
            prev_total = _total(prev_data[hid])  # type: ignore[index]
            curr_total = _total(data[hid])
            if prev_total is None or curr_total is None or prev_total == curr_total:
                continue
            diff = curr_total - prev_total
            r = data[hid]
            _write_row([
                "値上がり" if diff > 0 else "値下がり", label,
                r.get("building_name", ""), r.get("address", ""),
                r.get("layout", ""),        r.get("area", ""),
                r.get("station", ""),
                _fmt_yen(prev_total), _fmt_yen(curr_total),
                _fmt_yen(diff, show_sign=True),
                r.get("url", ""),
            ], UP_FILL if diff > 0 else DOWN_FILL)

        prev_data = data

    ws.freeze_panes = "A2"
    print(f"  変動イベント数: {row_n - 2} 件")


# ---------------------------------------------------------------------------
# Sheet 3: 合計賃料推移グラフ
# ---------------------------------------------------------------------------

def _build_chart_html(
    runs: list[tuple[str, dict[str, dict]]],
    output_path: Path,
) -> Path | None:
    """合計賃料推移の Plotly インタラクティブグラフを HTML として保存する。
    - 単発掲載（即入居）: 緑の点  ← 優秀
    - 複数回掲載（まだ残っている）: 赤の折れ線  ← 良くない
    """
    try:
        import plotly.graph_objects as go
        import plotly.colors as pcolors
    except ImportError:
        print("  plotly が見つかりません。スキップします。")
        return None

    all_ids: set[str] = set()
    for _, data in runs:
        all_ids.update(data.keys())

    if not all_ids:
        return None

    run_counts = {hid: sum(1 for _, data in runs if hid in data) for hid in all_ids}
    multi_ids  = sorted([hid for hid in all_ids if run_counts[hid] >= 2],
                        key=lambda h: -run_counts[h])
    single_ids = [hid for hid in all_ids if run_counts[hid] == 1]

    names: dict[str, str] = {}
    for hid in all_ids:
        for _, data in runs:
            if hid in data:
                names[hid] = data[hid].get("building_name", hid[:10])
                break

    # x 軸の値。runs は実行日時順に並んでいるので、その順序をそのまま軸に反映する。
    # 全ラベルが日時として解釈できれば日付軸（実際の間隔を反映）、
    # 1つでも解釈できなければカテゴリ軸に runs の順序を明示する。
    run_labels = [label for label, _ in runs]
    run_dts    = [_label_to_dt(label) for label in run_labels]
    use_date   = all(dt is not None for dt in run_dts)
    x_of: dict[str, object] = {
        label: (dt if use_date else label)
        for label, dt in zip(run_labels, run_dts)
    }

    def _total_man(row: dict) -> float | None:
        rent_val    = _parse_rent(row.get("rent", ""))
        kanrihi_val = _parse_kanrihi(row.get("kanrihi", ""))
        if rent_val is None:
            return None
        return round((rent_val + (kanrihi_val or 0)) / 10000, 2)

    # 掲載回数 → 色（黄→橙→赤、回数が多いほど濃い赤）
    multi_counts = [run_counts[hid] for hid in multi_ids] if multi_ids else [2]
    min_c, max_c = min(multi_counts), max(multi_counts)

    def _color_for_count(count: int) -> str:
        t = 0.0 if max_c <= min_c else (count - min_c) / (max_c - min_c)
        t = 0.25 + t * 0.75          # 0.25〜1.0 に収め薄すぎを回避
        return pcolors.sample_colorscale("YlOrRd", [t])[0]

    fig = go.Figure()

    # ---- 単発掲載: 1つのトレース（緑の点）----
    sx, sy, st = [], [], []
    for hid in single_ids:
        for label, data in runs:
            if hid in data:
                val = _total_man(data[hid])
                if val is not None:
                    sx.append(x_of[label])
                    sy.append(val)
                    st.append(names[hid])

    if sx:
        fig.add_trace(go.Scatter(
            x=sx, y=sy,
            mode="markers",
            marker=dict(color="#27ae60", size=10, opacity=0.8,
                        line=dict(color="#1a7a45", width=1)),
            name="単発掲載（即入居）",
            text=st,
            hovertemplate=(
                "<b>%{text}</b><br>"
                "実行日: %{x|%Y/%m/%d %H:%M}<br>"
                "合計賃料: %{y:.1f}万円"
                "<extra>単発掲載</extra>"
            ),
        ))

    # ---- 複数回掲載: 掲載回数ごとに色を変えた折れ線 ----
    for hid in multi_ids:
        xs, ys = [], []
        for label, data in runs:
            if hid in data:
                val = _total_man(data[hid])
                if val is not None:
                    xs.append(x_of[label])
                    ys.append(val)
        if not xs:
            continue
        count = run_counts[hid]
        color = _color_for_count(count)
        name  = names[hid]
        fig.add_trace(go.Scatter(
            x=xs, y=ys,
            mode="lines+markers",
            marker=dict(color=color, size=7, line=dict(color=color, width=1)),
            line=dict(color=color, width=2),
            name=f"{name}（{count}回）",
            hovertemplate=(
                f"<b>{name}</b><br>"
                f"掲載回数: {count}回<br>"
                "実行日: %{x|%Y/%m/%d %H:%M}<br>"
                "合計賃料: %{y:.1f}万円"
                f"<extra>{count}回掲載</extra>"
            ),
        ))

    # ---- カラースケールの凡例（ダミートレース） ----
    if multi_ids:
        fig.add_trace(go.Scatter(
            x=[None], y=[None],
            mode="markers",
            marker=dict(
                colorscale="YlOrRd",
                cmin=min_c, cmax=max_c,
                color=[min_c],
                colorbar=dict(
                    title=dict(text="掲載回数", side="right"),
                    thickness=14, len=0.5, x=1.02,
                    tickvals=list(range(min_c, max_c + 1)),
                ),
                showscale=True, size=0,
            ),
            showlegend=False,
            hoverinfo="skip",
        ))

    fig.update_layout(
        title=dict(text="合計賃料推移（賃料＋管理費）", font=dict(size=18)),
        autosize=True,          # ブラウザウィンドウに追従
        margin=dict(l=70, r=130, t=60, b=90),
        xaxis=dict(
            title="実行日",
            tickangle=-30,
            showgrid=True,
            gridcolor="#eeeeee",
            **(dict(type="date", tickformat="%Y/%m/%d<br>%H:%M")
               if use_date else
               dict(type="category", categoryorder="array", categoryarray=run_labels)),
        ),
        yaxis=dict(
            title="合計賃料（万円）",
            tickformat=".1f",
            ticksuffix="万円",
            showgrid=True,
            gridcolor="#eeeeee",
            zeroline=False,
        ),
        hovermode="closest",
        legend=dict(
            title=dict(text="物件（凡例クリックで表示切替）"),
            bgcolor="rgba(255,255,255,0.85)",
            bordercolor="#cccccc",
            borderwidth=1,
        ),
        plot_bgcolor="white",
        paper_bgcolor="white",
    )

    # ビューポート全体に広げる CSS を注入して保存
    html = fig.to_html(
        include_plotlyjs=True,
        full_html=True,
        config={"responsive": True},
    )
    css = (
        "html,body{margin:0;padding:0;height:100%;overflow:hidden}"
        ".plotly-graph-div{height:100vh!important;width:100vw!important}"
    )
    html = html.replace("</head>", f"<style>{css}</style></head>")

    html_path = output_path.with_name(output_path.stem + "_chart.html")
    html_path.write_text(html, encoding="utf-8")
    return html_path


def _build_chart_sheet(wb: Workbook, runs: list[tuple[str, dict[str, dict]]],
                       output_path: Path) -> None:
    """グラフ HTML を生成し、Excel に参照先を記載する。"""
    html_path = _build_chart_html(runs, output_path)

    ws = wb.create_sheet("合計賃料推移")
    if html_path:
        ws["A1"] = f"インタラクティブグラフ → {html_path.name}"
        ws["A2"] = "（同じフォルダの上記 HTML ファイルをブラウザで開いてください）"
        ws["A3"] = "凡例クリック: 物件を表示/非表示   ホバー: 物件名・賃料を確認"
        ws.column_dimensions["A"].width = 60
        print(f"  グラフ HTML: {html_path.name}")


# ---------------------------------------------------------------------------
# メイン
# ---------------------------------------------------------------------------

def compare(run_dirs: list[Path], output_path: Path) -> None:
    runs: list[tuple[str, dict[str, dict]]] = []
    for rd in run_dirs:
        data = _load_run(rd)
        runs.append((_run_label(rd), data))
        print(f"読み込み: {rd.name} ({len(data)} 件)")

    wb = Workbook()

    print("\n[Sheet 1] 時系列推移を生成中...")
    _build_matrix_sheet(wb, runs, run_dirs)

    print("[Sheet 2] 変動サマリーを生成中...")
    _build_diff_sheet(wb, runs)

    print("[Sheet 3] 合計賃料推移グラフを生成中...")
    _build_chart_sheet(wb, runs, output_path)

    wb.save(str(output_path))
    print(f"\n保存: {output_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="物件の時系列推移を比較する Excel を生成")
    parser.add_argument(
        "--last", type=int, default=None, metavar="N",
        help="直近 N 回分のみ対象にする（省略時は全件）",
    )
    parser.add_argument(
        "--output", default=None, metavar="PATH",
        help="出力ファイルパス（省略時: output/comparison.xlsx）",
    )
    args = parser.parse_args()

    run_dirs = _find_all_run_dirs(OUTPUT_ROOT)
    if not run_dirs:
        print("results.csv を持つ run_dir が見つかりません。", file=sys.stderr)
        sys.exit(1)

    if args.last:
        run_dirs = run_dirs[-args.last:]

    print(f"対象: {len(run_dirs)} 回分 ({run_dirs[0].name} 〜 {run_dirs[-1].name})")

    output_path = Path(args.output) if args.output else OUTPUT_ROOT / "comparison.xlsx"
    compare(run_dirs, output_path)


if __name__ == "__main__":
    main()
